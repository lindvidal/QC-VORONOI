"""
GitHub Commit Logger & Telemetry Service
========================================
Exports git commit extraction logs directly to Docs without generating intermediate scripts.
Always generates log_github_lastrun.csv and log_github_lastrun.docx in src/docs/telemetry/.
"""

import os
import subprocess
import csv
from pathlib import Path
from docx import Document

SCRIPT_DIR = Path(__file__).resolve().parent
PROJECT_ROOT = SCRIPT_DIR.parent.parent
TELEMETRY_DIR = SCRIPT_DIR.parent / "docs" / "telemetry"


def generate_git_log():
    TELEMETRY_DIR.mkdir(parents=True, exist_ok=True)
    csv_path = TELEMETRY_DIR / "log_github_lastrun.csv"
    docx_path = TELEMETRY_DIR / "log_github_lastrun.docx"

    # 1. Run git log command directly
    cmd = ['git', 'log', '--pretty=format:%h|%ad|%an|%s', '--date=short']
    res = subprocess.run(cmd, cwd=PROJECT_ROOT, capture_output=True, text=True, encoding="utf-8")
    
    lines = res.stdout.strip().split("\n") if res.stdout else []
    
    # Write CSV
    with open(csv_path, "w", encoding="utf-8", newline="") as f:
        writer = csv.writer(f, delimiter="|")
        writer.writerow(["Hash", "Date", "Author", "Message"])
        for line in lines:
            if line.strip():
                parts = line.split("|", 3)
                if len(parts) == 4:
                    writer.writerow(parts)

    # 2. Convert to DOCX
    doc = Document()
    doc.add_heading("GitHub Commit Log - QC-VORONOI", 0)

    try:
        with open(csv_path, "r", encoding="utf-8") as f:
            reader = csv.reader(f, delimiter="|")
            rows = list(reader)
            if rows:
                table = doc.add_table(rows=len(rows), cols=4)
                table.style = "Table Grid"
                for i, row in enumerate(rows):
                    for j, val in enumerate(row):
                        table.rows[i].cells[j].text = val
    except Exception as e:
        doc.add_paragraph(f"Error reading CSV: {e}")

    doc.save(str(docx_path))
    print(f"[OK] Telemetry logs saved:\n  - {csv_path}\n  - {docx_path}")


if __name__ == "__main__":
    generate_git_log()
